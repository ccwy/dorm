"""
iOS 专用 pandas 兼容层（基于 openpyxl / xlsxwriter）

iOS 端 kivy-ios 编译的 Python.framework 无法编译 pandas 的 C 扩展，
本模块用纯 Python 实现 pandas 的一个 API 子集，覆盖项目中导入导出
文件实际使用的功能，使现有代码零改动运行。

覆盖的 API：
- 模块级：isna / notna / Timestamp / NA / NaN / read_excel / ExcelWriter / ExcelFile
- Series：tolist / fillna / __iter__ / __len__
- DataFrame：columns / empty / iterrows / rename / get / __getitem__ /
  to_excel / index / sort_index / loc 赋值

依赖：openpyxl（必装）、xlsxwriter（engine='xlsxwriter' 时需要）。
注意：openpyxl / xlsxwriter 均为延迟导入，不在模块顶层 import，
避免模块加载失败影响 stub 安装。
"""

import datetime
import io as _io

# ==================== 模块级常量与函数 ====================

NA = float('nan')
NaN = float('nan')

# pandas.Timestamp 在项目中仅用于 isinstance 检查（excel_date_utils），
# datetime 实例会在更早的分支返回，此处别名只为兼容该检查
Timestamp = datetime.datetime


def isna(v):
    """判断值是否为缺失（None 或 NaN）"""
    return v is None or (isinstance(v, float) and v != v)


def notna(v):
    """判断值是否非缺失"""
    return not isna(v)


# ==================== Series ====================


class Series:
    """一维数据，内部用 list 存储"""

    def __init__(self, data=None):
        self._data = list(data) if data is not None else []

    def tolist(self):
        return list(self._data)

    def fillna(self, value):
        return Series(value if isna(v) else v for v in self._data)

    def __iter__(self):
        return iter(self._data)

    def __len__(self):
        return len(self._data)

    def __getitem__(self, i):
        return self._data[i]

    def __repr__(self):
        return f'Series({self._data!r})'


# ==================== Index 辅助类 ====================


class _Columns:
    """模拟 pd.Index（列名集合）"""

    def __init__(self, cols):
        self._cols = list(cols)

    def __iter__(self):
        return iter(self._cols)

    def __contains__(self, col):
        return col in self._cols

    def tolist(self):
        return list(self._cols)

    @property
    def values(self):
        return list(self._cols)

    def __repr__(self):
        return f'Index({self._cols!r})'


class _Index:
    """模拟 pd.RangeIndex，支持 df.index + int"""

    def __init__(self, labels):
        self._labels = list(labels)

    def __add__(self, n):
        return _Index(x + n for x in self._labels)

    def __iter__(self):
        return iter(self._labels)

    def __len__(self):
        return len(self._labels)

    def __getitem__(self, i):
        return self._labels[i]

    def __repr__(self):
        return f'RangeIndex({self._labels!r})'


# ==================== DataFrame ====================


class _LocAccessor:
    """支持 df.loc[label] = row 赋值（行追加）"""

    def __init__(self, df):
        self._df = df

    def __setitem__(self, label, row_values):
        if isinstance(row_values, dict):
            row = dict(row_values)
        else:
            row = dict(zip(self._df._columns._cols, row_values))
        self._df._rows.append(row)
        self._df._index._labels.append(label)


class DataFrame:
    """二维表，内部用 list[dict] 行式存储"""

    def __init__(self, data=None, columns=None):
        if isinstance(data, dict):
            # 列式 dict[list] → 行式
            cols = list(data.keys())
            n = max((len(v) for v in data.values()), default=0)
            self._rows = [
                {c: (data[c][i] if i < len(data[c]) else None) for c in cols}
                for i in range(n)
            ]
            self._columns = _Columns(cols)
        elif isinstance(data, list) and data and isinstance(data[0], dict):
            # 行式 list[dict]，指定 columns 时按其取值与排序
            cols = list(columns) if columns is not None else list(data[0].keys())
            self._columns = _Columns(cols)
            self._rows = [{c: r.get(c) for c in cols} for r in data]
        elif isinstance(data, list) and columns is not None:
            # list[list] + 列名
            self._columns = _Columns(list(columns))
            self._rows = [dict(zip(self._columns._cols, row)) for row in data]
        else:
            # 空表（data 为 None 或空 list）
            self._columns = _Columns(list(columns) if columns else [])
            self._rows = []
        self._index = _Index(range(len(self._rows)))

    @property
    def columns(self):
        return self._columns

    @property
    def empty(self):
        return len(self._rows) == 0

    @property
    def index(self):
        return self._index

    @index.setter
    def index(self, value):
        self._index = _Index(value)

    @property
    def loc(self):
        return _LocAccessor(self)

    def __len__(self):
        return len(self._rows)

    def __getitem__(self, col):
        return Series(row.get(col) for row in self._rows)

    def get(self, col, default=None):
        if col in self._columns:
            return self.__getitem__(col)
        return default

    def iterrows(self):
        for label, row in zip(self._index._labels, self._rows):
            yield label, row

    def rename(self, columns=None):
        mapping = columns or {}
        new_rows = [{mapping.get(k, k): v for k, v in row.items()} for row in self._rows]
        return DataFrame(new_rows)

    def sort_index(self):
        pairs = sorted(zip(self._index._labels, self._rows), key=lambda p: p[0])
        return DataFrame([r for _, r in pairs], columns=self._columns._cols)

    def to_excel(self, excel_writer, sheet_name='Sheet1', index=False,
                 startrow=0, header=True, **kwargs):
        if isinstance(excel_writer, ExcelWriter):
            excel_writer._write_dataframe(self, sheet_name, startrow, header)
        else:
            # 传入路径：自建 writer 并保存
            writer = ExcelWriter(excel_writer)
            try:
                writer._write_dataframe(self, sheet_name, startrow, header)
            finally:
                writer.close()


# ==================== ExcelWriter ====================


class ExcelWriter:
    """pandas.ExcelWriter 兼容层，支持 openpyxl / xlsxwriter 引擎"""

    def __init__(self, path, engine=None):
        self._path = path
        self._engine = engine or 'openpyxl'
        self.sheets = {}
        if self._engine == 'openpyxl':
            from openpyxl import Workbook
            self.book = Workbook()
            self.book.remove(self.book.active)
        elif self._engine == 'xlsxwriter':
            import xlsxwriter
            self.book = xlsxwriter.Workbook(path, {'in_memory': True})
        else:
            raise ValueError(f"不支持的 ExcelWriter 引擎: {self._engine}")

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()
        return False

    def close(self):
        if self._engine == 'openpyxl':
            self.book.save(self._path)
        else:
            self.book.close()

    def _write_dataframe(self, df, sheet_name, startrow, header):
        cols = df.columns._cols
        if self._engine == 'openpyxl':
            ws = self.book.create_sheet(title=sheet_name)
            r = startrow + 1  # openpyxl 行号 1-based
            if header:
                for c, name in enumerate(cols, 1):
                    ws.cell(row=r, column=c, value=name)
                r += 1
            for row in df._rows:
                for c, name in enumerate(cols, 1):
                    ws.cell(row=r, column=c, value=row.get(name))
                r += 1
            self.sheets[sheet_name] = ws
        else:
            ws = self.book.add_worksheet(sheet_name)
            r = startrow  # xlsxwriter 行号 0-based
            if header:
                for c, name in enumerate(cols):
                    ws.write(r, c, name)
                r += 1
            for row in df._rows:
                for c, name in enumerate(cols):
                    ws.write(r, c, row.get(name))
                r += 1
            self.sheets[sheet_name] = ws


# ==================== ExcelFile / read_excel ====================


class ExcelFile:
    """pandas.ExcelFile 兼容层，延迟加载 openpyxl"""

    def __init__(self, io):
        from openpyxl import load_workbook
        if isinstance(io, bytes):
            io = _io.BytesIO(io)
        self.book = load_workbook(io, read_only=True, data_only=True)

    @property
    def sheet_names(self):
        return self.book.sheetnames


def read_excel(io, sheet_name=0, nrows=None, dtype=None, **kwargs):
    """读取 Excel 为 DataFrame，第一行为列名"""
    from openpyxl import load_workbook

    if isinstance(io, ExcelFile):
        wb = io.book
    else:
        if isinstance(io, bytes):
            io = _io.BytesIO(io)
        wb = load_workbook(io, read_only=True, data_only=True)

    ws = wb[sheet_name] if isinstance(sheet_name, str) else wb.worksheets[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return DataFrame()

    headers = list(rows[0])
    data_rows = rows[1:]
    if nrows is not None:
        data_rows = data_rows[:nrows]

    dicts = [
        {h: (row[i] if i < len(row) else None) for i, h in enumerate(headers)}
        for row in data_rows
    ]
    df = DataFrame(dicts)

    # dtype: {列名: 类型函数}，对对应列每个非空值调用
    if dtype:
        for col, conv in dtype.items():
            if col in df._columns:
                for row in df._rows:
                    v = row.get(col)
                    if notna(v):
                        row[col] = conv(v)
    return df